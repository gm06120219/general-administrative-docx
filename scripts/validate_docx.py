#!/usr/bin/env python3
"""Validate the main structural formatting rules of a generated DOCX."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn


EXPECTED_STYLES = {
    "Normal": (16, "仿宋_GB2312"),
    "Title": (22, "方正小标宋简体"),
    "Heading 1": (16, "黑体"),
    "Heading 2": (16, "楷体_GB2312"),
    "Heading 3": (16, "仿宋_GB2312"),
    "Heading 4": (16, "仿宋_GB2312"),
}


def close(actual: float, expected: float, tolerance: float = 0.08) -> bool:
    return abs(actual - expected) <= tolerance


def style_font(style) -> Optional[str]:
    rpr = style.element.rPr
    if rpr is None or rpr.rFonts is None:
        return style.font.name
    return rpr.rFonts.get(qn("w:eastAsia")) or style.font.name


def validate(path: Path) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    doc = Document(path)
    if not doc.sections:
        errors.append("文档没有节定义")
        return {"ok": False, "errors": errors, "warnings": warnings}

    section = doc.sections[0]
    checks = {
        "page_width_cm": (section.page_width.cm, 21.0),
        "page_height_cm": (section.page_height.cm, 29.7),
        "top_margin_cm": (section.top_margin.cm, 3.7),
        "bottom_margin_cm": (section.bottom_margin.cm, 3.7),
        "left_margin_cm": (section.left_margin.cm, 2.8),
        "right_margin_cm": (section.right_margin.cm, 2.8),
    }
    for name, (actual, expected) in checks.items():
        if not close(actual, expected):
            errors.append(f"{name}: expected {expected}, got {actual:.3f}")

    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is None or grid.get(qn("w:linePitch")) != "560":
        errors.append("缺少 28 pt 文档网格（w:docGrid linePitch=560）")

    for name, (expected_size, expected_font) in EXPECTED_STYLES.items():
        if name not in doc.styles:
            errors.append(f"缺少样式：{name}")
            continue
        style = doc.styles[name]
        actual_size = style.font.size.pt if style.font.size else None
        actual_font = style_font(style)
        if actual_size is None or not close(actual_size, expected_size, 0.1):
            errors.append(f"{name} 字号应为 {expected_size} pt，实际为 {actual_size}")
        if actual_font != expected_font:
            warnings.append(f"{name} 字体为 {actual_font!r}，默认期望 {expected_font!r}（可能是显式字体覆盖）")

    normal = doc.styles["Normal"].paragraph_format
    if normal.line_spacing_rule != WD_LINE_SPACING.EXACTLY:
        errors.append("正文样式不是固定行距")
    if normal.line_spacing is None or not close(normal.line_spacing.pt, 28, 0.1):
        errors.append("正文固定行距不是 28 pt")

    title_ppr = doc.styles["Title"].element.pPr
    if title_ppr is not None and title_ppr.find(qn("w:pBdr")) is not None:
        errors.append("标题样式含段落边框或装饰线")

    heading_checks = {
        "Heading 1": re.compile(r"^[一二三四五六七八九十百]+、"),
        "Heading 2": re.compile(r"^（[一二三四五六七八九十百]+）"),
        "Heading 3": re.compile(r"^\d+．"),
        "Heading 4": re.compile(r"^（\d+）"),
    }
    for paragraph in doc.paragraphs:
        if paragraph.style.name in heading_checks and paragraph.text:
            if not heading_checks[paragraph.style.name].search(paragraph.text):
                warnings.append(f"标题序号格式可疑：{paragraph.text}")

    if not re.search(r"-\d{8}-V\d+\.\d+\.docx$", path.name):
        warnings.append("文件名未采用 标题-YYYYMMDD-VN.N.docx")
    if not doc.paragraphs or not doc.paragraphs[0].text.strip():
        errors.append("文档标题为空")

    return {
        "ok": not errors,
        "file": str(path),
        "paragraphs": len(doc.paragraphs),
        "errors": errors,
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", help="DOCX file to validate")
    parser.add_argument("--json", action="store_true", help="emit machine-readable JSON")
    args = parser.parse_args()
    path = Path(args.docx).resolve()
    if not path.exists():
        print(f"ERROR: file does not exist: {path}", file=sys.stderr)
        return 2
    report = validate(path)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print("PASS" if report["ok"] else "FAIL", report["file"])
        for item in report["errors"]:
            print("ERROR:", item)
        for item in report["warnings"]:
            print("WARN:", item)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
