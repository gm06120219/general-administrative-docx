#!/usr/bin/env python3
"""Generate a formatted Chinese administrative DOCX from a JSON specification."""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path
from typing import Any, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DEFAULT_FONTS = {
    "title": "方正小标宋简体",
    "body": "仿宋_GB2312",
    "h1": "黑体",
    "h2": "楷体_GB2312",
    "h3": "仿宋_GB2312",
    "h4": "仿宋_GB2312",
    "page_number": "宋体",
}

SIZE_BODY = 16
SIZE_TITLE = 22
LINE_SPACING = 28
FULLWIDTH_SPACE = "　"


def set_font(element: Any, family: str, size_pt: float, bold: Optional[bool] = None) -> None:
    """Set Latin and East Asian font properties on a style or run."""
    element.font.name = family
    element.font.size = Pt(size_pt)
    element.font.color.rgb = RGBColor(0, 0, 0)
    element.font.underline = False
    element.font.italic = False
    if bold is not None:
        element.font.bold = bold
    rpr = element._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), family)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_paragraph_grid(paragraph: Any, *, keep_with_next: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(LINE_SPACING)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = keep_with_next
    ppr = paragraph._p.get_or_add_pPr()
    snap = ppr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        ppr.append(snap)
    snap.set(qn("w:val"), "1")
    widow = ppr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        ppr.append(widow)
    widow.set(qn("w:val"), "0")


def configure_style(doc: Document, name: str, font: str, size: float, *, bold: bool = False,
                    align: Optional[WD_ALIGN_PARAGRAPH] = None, first_indent_chars: float = 0,
                    keep_with_next: bool = False) -> Any:
    styles = doc.styles
    if name in styles:
        style = styles[name]
    else:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_font(style, font, size, bold)
    fmt = style.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(LINE_SPACING)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(size * first_indent_chars) if first_indent_chars else Pt(0)
    fmt.keep_with_next = keep_with_next
    if align is not None:
        fmt.alignment = align
    return style


def remove_paragraph_borders(paragraph_or_style: Any) -> None:
    ppr = paragraph_or_style._element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def configure_document(doc: Document, fonts: dict[str, str]) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.7)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(3.0)
    section.start_type = WD_SECTION.NEW_PAGE

    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:docGrid")):
        sect_pr.remove(old)
    grid = OxmlElement("w:docGrid")
    grid.set(qn("w:type"), "linesAndChars")
    grid.set(qn("w:linePitch"), str(LINE_SPACING * 20))
    grid.set(qn("w:charSpace"), "0")
    sect_pr.append(grid)

    configure_style(doc, "Normal", fonts["body"], SIZE_BODY, first_indent_chars=2)
    title = configure_style(
        doc, "Title", fonts["title"], SIZE_TITLE, align=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True
    )
    remove_paragraph_borders(title)
    title.paragraph_format.space_after = Pt(LINE_SPACING)
    configure_style(doc, "Heading 1", fonts["h1"], SIZE_BODY, keep_with_next=True)
    configure_style(doc, "Heading 2", fonts["h2"], SIZE_BODY, keep_with_next=True)
    configure_style(doc, "Heading 3", fonts["h3"], SIZE_BODY, bold=True, keep_with_next=True)
    configure_style(doc, "Heading 4", fonts["h4"], SIZE_BODY, keep_with_next=True)

    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    setting = OxmlElement("w:compatSetting")
    setting.set(qn("w:name"), "compatibilityMode")
    setting.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    setting.set(qn("w:val"), "15")
    compat.append(setting)


def chinese_number(number: int) -> str:
    digits = "零一二三四五六七八九"
    if number < 10:
        return digits[number]
    if number == 10:
        return "十"
    if number < 20:
        return "十" + digits[number % 10]
    if number < 100:
        tens, ones = divmod(number, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(number)


PREFIX_PATTERNS = {
    1: re.compile(r"^\s*[一二三四五六七八九十百]+、"),
    2: re.compile(r"^\s*（[一二三四五六七八九十百]+）"),
    3: re.compile(r"^\s*\d+[．.]"),
    4: re.compile(r"^\s*[（(]\d+[）)]"),
}


def add_heading_number(text: str, level: int, counters: list[int]) -> str:
    if PREFIX_PATTERNS[level].search(text):
        return text
    n = counters[level]
    if level == 1:
        return f"{chinese_number(n)}、{text}"
    if level == 2:
        return f"（{chinese_number(n)}）{text}"
    if level == 3:
        return f"{n}．{text}"
    return f"（{n}）{text}"


def normalize_two_char_name(name: str) -> str:
    value = name.strip().replace(FULLWIDTH_SPACE, "")
    if len(value) == 2 and all("\u3400" <= ch <= "\u9fff" for ch in value):
        return value[0] + FULLWIDTH_SPACE + value[1]
    return name.strip()


def apply_direct_font(run: Any, family: str, size: float, bold: bool | None = None) -> None:
    set_font(run, family, size, bold)


def add_page_field(paragraph: Any, font: str) -> None:
    paragraph.add_run("— ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    paragraph.add_run(" —")
    for item in paragraph.runs:
        apply_direct_font(item, font, 14, False)


def configure_page_numbers(doc: Document, font: str) -> None:
    settings = doc.settings._element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    for section in doc.sections:
        odd = section.footer.paragraphs[0]
        odd.clear()
        odd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        odd.paragraph_format.right_indent = Pt(SIZE_BODY)
        set_paragraph_grid(odd)
        add_page_field(odd, font)

        even = section.even_page_footer.paragraphs[0]
        even.clear()
        even.alignment = WD_ALIGN_PARAGRAPH.LEFT
        even.paragraph_format.left_indent = Pt(SIZE_BODY)
        set_paragraph_grid(even)
        add_page_field(even, font)


def add_paragraph(doc: Document, text: str, font: str, *, first_indent_chars: float = 2,
                  bold: bool = False, align: Optional[WD_ALIGN_PARAGRAPH] = None) -> Any:
    paragraph = doc.add_paragraph(style="Normal")
    set_paragraph_grid(paragraph)
    paragraph.paragraph_format.first_line_indent = Pt(SIZE_BODY * first_indent_chars)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    apply_direct_font(run, font, SIZE_BODY, bold)
    return paragraph


def build_document(data: dict[str, Any]) -> Document:
    if not data.get("title"):
        raise ValueError("input JSON must contain a non-empty 'title'")
    if not isinstance(data.get("blocks"), list):
        raise ValueError("input JSON must contain a 'blocks' array")

    fonts = dict(DEFAULT_FONTS)
    fonts.update({k: v for k, v in (data.get("fonts") or {}).items() if k in fonts and v})

    doc = Document()
    configure_document(doc, fonts)
    doc.core_properties.title = str(data["title"])
    doc.core_properties.subject = str(data.get("document_type", "行政公文"))
    doc.core_properties.author = str(data.get("unit_name", ""))

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    set_paragraph_grid(title, keep_with_next=True)
    title.paragraph_format.space_after = Pt(LINE_SPACING)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(str(data["title"]))
    apply_direct_font(title_run, fonts["title"], SIZE_TITLE, False)

    if data.get("recipient"):
        add_paragraph(doc, str(data["recipient"]), fonts["body"], first_indent_chars=0)

    counters = [0, 0, 0, 0, 0]
    for block in data["blocks"]:
        if not isinstance(block, dict):
            raise ValueError("every block must be an object")
        kind = str(block.get("type", "paragraph"))
        text = str(block.get("text", ""))

        if kind.startswith("heading") and kind[-1:].isdigit():
            level = int(kind[-1])
            if level not in (1, 2, 3, 4):
                raise ValueError(f"unsupported heading level: {level}")
            counters[level] += 1
            for lower in range(level + 1, 5):
                counters[lower] = 0
            if block.get("auto_number", True):
                text = add_heading_number(text, level, counters)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            set_paragraph_grid(paragraph, keep_with_next=True)
            paragraph.paragraph_format.first_line_indent = Pt(0)
            run = paragraph.add_run(text)
            apply_direct_font(run, fonts[f"h{level}"], SIZE_BODY, level == 3)
            continue

        if kind == "paragraph":
            add_paragraph(
                doc,
                text,
                fonts["body"],
                first_indent_chars=float(block.get("first_line_indent_chars", 2)),
                bold=bool(block.get("bold", False)),
            )
        elif kind == "names":
            names = [normalize_two_char_name(str(name)) for name in block.get("names", [])]
            per_line = max(1, int(block.get("per_line", 5)))
            for start in range(0, len(names), per_line):
                add_paragraph(
                    doc,
                    FULLWIDTH_SPACE.join(names[start:start + per_line]),
                    fonts["body"],
                    first_indent_chars=0,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
        elif kind == "attachment":
            value = text if text.startswith("附件") else f"附件：{text}"
            add_paragraph(doc, value, fonts["body"], first_indent_chars=0)
        elif kind == "page_break":
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        elif kind == "blank":
            add_paragraph(doc, "", fonts["body"], first_indent_chars=0)
        else:
            raise ValueError(f"unsupported block type: {kind}")

    signature = data.get("signature", None)
    if signature is not False:
        if signature is None:
            signature = {}
        unit_name = str(signature.get("unit_name") or data.get("unit_name") or "").strip()
        date_text = str(signature.get("date") or data.get("date") or chinese_date(date.today())).strip()
        for value in (unit_name, date_text):
            if not value:
                continue
            paragraph = add_paragraph(
                doc, value, fonts["body"], first_indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.left_indent = Pt(SIZE_BODY * 10)

    mode = str(data.get("mode", "ordinary")).lower()
    page_numbers = data.get("page_numbers")
    if page_numbers is None:
        page_numbers = mode == "redhead"
    if page_numbers:
        configure_page_numbers(doc, fonts["page_number"])
    return doc


def chinese_date(value: date) -> str:
    return f"{value.year}年{value.month}月{value.day}日"


def safe_stem(value: str) -> str:
    value = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", value).strip().rstrip(".")
    return value[:100] or "行政公文"


def normalized_version(value: str) -> Tuple[int, int]:
    match = re.fullmatch(r"[Vv]?(\d+)(?:\.(\d+))?", value.strip())
    if not match:
        raise ValueError("version must look like 1.0 or V1.0")
    return int(match.group(1)), int(match.group(2) or 0)


def resolve_output(data: dict[str, Any], output_dir: Path, output: Optional[str],
                   file_date: Optional[str], overwrite: bool) -> Path:
    if output:
        path = Path(output)
        if not path.is_absolute():
            path = output_dir / path
        if path.suffix.lower() != ".docx":
            path = path.with_suffix(".docx")
        if path.exists() and not overwrite:
            raise FileExistsError(f"output already exists: {path}")
        return path

    stamp = file_date or date.today().strftime("%Y%m%d")
    if not re.fullmatch(r"\d{8}", stamp):
        raise ValueError("file date must use YYYYMMDD")
    major, minor = normalized_version(str(data.get("version", "1.0")))
    stem = safe_stem(str(data["title"]))
    while True:
        candidate = output_dir / f"{stem}-{stamp}-V{major}.{minor}.docx"
        if overwrite or not candidate.exists():
            return candidate
        minor += 1


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="UTF-8 JSON input file")
    parser.add_argument("--output-dir", default=".", help="directory for the generated DOCX")
    parser.add_argument("--output", help="explicit output filename or path")
    parser.add_argument("--file-date", help="filename date in YYYYMMDD; defaults to today")
    parser.add_argument("--overwrite", action="store_true", help="allow replacing an explicit/default path")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        data = json.loads(input_path.read_text(encoding="utf-8"))
        doc = build_document(data)
        destination = resolve_output(data, output_dir, args.output, args.file_date, args.overwrite)
        destination.parent.mkdir(parents=True, exist_ok=True)
        doc.save(destination)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
